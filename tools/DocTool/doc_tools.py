"""DocTool.doc_tools — document read/create tools (extracted from tools.py).

Self-contained: only ``os`` at module level; heavy libs (pymupdf, python-docx,
openpyxl) are imported lazily inside each function so the app boots without
them. tools.DocTool re-exports these so tools.py imports them by name.
"""
import os


def read_pdf(file_path: str, max_pages: int = 20, page_numbers: str = "") -> str:
    """Read text from a PDF file.

    Args:
        file_path: Path to the PDF file.
        max_pages: Max pages to read (default 20). Use 0 for all.
        page_numbers: Specific pages to read, e.g. '1-3,5,7-9' (overrides max_pages).
    """
    if not os.path.isfile(file_path):
        return "Error: File not found: " + file_path
    if not file_path.lower().endswith('.pdf'):
        return "Error: Not a PDF file: " + file_path
    try:
        import fitz
    except ImportError:
        return "Error: pymupdf not installed. Run: pip install pymupdf"
    try:
        doc = fitz.open(file_path)
        total = len(doc)
        pages = []
        if page_numbers:
            import re
            for part in page_numbers.split(','):
                part = part.strip()
                if '-' in part:
                    a, b = part.split('-', 1)
                    pages.extend(range(int(a.strip())-1, int(b.strip())))
                else:
                    pages.append(int(part)-1)
        else:
            max_p = max_pages if max_pages > 0 else total
            pages = list(range(min(total, max_p)))
        output = []
        output.append("PDF: " + os.path.basename(file_path) + " (" + str(total) + " pages)")
        for pnum in pages:
            if pnum >= total:
                break
            page = doc[pnum]
            text = page.get_text().strip()
            if text:
                output.append("--- Page " + str(pnum+1) + " ---")
                # Truncate very long pages to 3000 chars
                if len(text) > 3000:
                    text = text[:1500] + "\n...[truncated]...\n" + text[-1500:]
                output.append(text)
        doc.close()
        result = "\n".join(output)
        if len(result) > 50000:
            result = result[:25000] + "\n...[large PDF truncated]...\n" + result[-25000:]
        return result if result.strip() else "(empty PDF)"
    except Exception as e:
        return "Error reading PDF: " + str(e)


def create_docx(file_path, title="", content=""):
    """Create a Word document (.docx) with optional title and content."""
    try:
        from docx import Document
    except ImportError:
        return "Error: python-docx not installed. Run: pip install python-docx"
    try:
        doc = Document()
        if title:
            doc.add_heading(title, level=1)
        for para in content.split(chr(10)):
            para = para.strip()
            if para:
                doc.add_paragraph(para)
        os.makedirs(os.path.dirname(file_path) or ".", exist_ok=True)
        doc.save(file_path)
        return "Created Word document: " + file_path
    except Exception as e:
        return "Error creating document: " + str(e)


def create_xlsx(file_path, headers=None, rows=None):
    """Create an Excel spreadsheet (.xlsx)."""
    try:
        from openpyxl import Workbook
    except ImportError:
        return "Error: openpyxl not installed. Run: pip install openpyxl"
    try:
        wb = Workbook()
        ws = wb.active
        headers = headers or []
        rows = rows or []
        if headers:
            ws.append(headers)
        for row in rows:
            ws.append(row if isinstance(row, list) else [row])
        os.makedirs(os.path.dirname(file_path) or ".", exist_ok=True)
        wb.save(file_path)
        return "Created Excel: " + file_path + " (" + str(len(headers)) + " cols, " + str(len(rows)) + " rows)"
    except Exception as e:
        return "Error creating spreadsheet: " + str(e)